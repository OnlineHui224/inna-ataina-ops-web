from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
import re
from src.models.data_schemas import TicketData
from src.core.logger import logger

class DocxGenerator:
    def __init__(self, template_path: str, output_dir: str):
        self.template_path = template_path
        self.output_dir = output_dir

    def _safe_write_cell(self, cell, text):
        """Writes standard text to a cell (used for flights and pax)."""
        if not cell: return
        cell.text = str(text)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = 'Arial'

    def _write_bold_centered_cell(self, cell, text):
        """Specifically formats hotel and group names to be Bold and Centered."""
        if not cell: return
        cell.text = "" # Clear the cell safely
        paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(str(text))
        run.font.name = 'Arial'
        run.bold = True

    def _format_flight_number(self, carrier: str, flight_no: str) -> str:
        """Formats flight numbers to strictly be CARRIER + 4 DIGITS (e.g., QR0635)"""
        carrier_code = str(carrier).strip().upper()
        digits_only = re.sub(r'\D', '', str(flight_no))
        
        if digits_only:
            padded_flight = digits_only.zfill(4)
            return f"{carrier_code}{padded_flight}"
        return f"{carrier_code}{str(flight_no).strip()}"

    def generate(self, data: TicketData, makkah_hotel: str = "", madinah_hotel: str = "", group_name: str = "") -> str:
        logger.info("Starting DOCX generation with targeted cell injection...")
        if not os.path.exists(self.template_path):
            raise FileNotFoundError(f"Template not found at {self.template_path}")
            
        if not os.path.exists(self.output_dir):
            os.makedirs(self.output_dir)

        doc = Document(self.template_path)

        # 1. Header Injection 
        for p in doc.paragraphs:
            if "FLIGHT:" in p.text and "MEDINAH" in p.text:
                p.text = f"FLIGHT: {data.primary_carrier.upper()} – MEDINAH"

        # 2. Table 1: Passenger Totals & Group Name
        # We use laser-targeted grid coordinates here (Row 2 in Python is technically the 3rd row visually)
        if len(doc.tables) > 0:
            table_pax = doc.tables[0]
            if len(table_pax.rows) > 2:
                # Target Column 1 (Group Name)
                if group_name:
                    self._write_bold_centered_cell(table_pax.cell(2, 1), group_name)
                    
                # Target Columns 2, 3, and 4 (Adult, Child, Total)
                self._safe_write_cell(table_pax.cell(2, 2), data.adults)   
                self._safe_write_cell(table_pax.cell(2, 3), data.children)  
                self._safe_write_cell(table_pax.cell(2, 4), data.total_pax) 

        # 3. Table 2: Flights
        if len(doc.tables) > 1:
            table_flights = doc.tables[1]
            start_row = 2 
            
            for idx, flight in enumerate(data.flights):
                row_idx = start_row + idx
                if row_idx >= len(table_flights.rows):
                    table_flights.add_row()
                
                row = table_flights.rows[row_idx]
                self._safe_write_cell(row.cells[0], flight.departure_city)
                self._safe_write_cell(row.cells[1], flight.arrival_city)
                
                date_str = f"{flight.date} 2026" if "2026" not in flight.date else flight.date
                self._safe_write_cell(row.cells[2], date_str)
                
                self._safe_write_cell(row.cells[3], flight.departure_time)
                self._safe_write_cell(row.cells[4], flight.arrival_time)
                self._safe_write_cell(row.cells[5], flight.carrier)
                
                perfect_flight_no = self._format_flight_number(flight.carrier, flight.flight_number)
                self._safe_write_cell(row.cells[6], perfect_flight_no)
                
                if idx == 0:
                    self._safe_write_cell(row.cells[7], data.pnr)

        # 4. Table 3: Hotel Accommodation
        for table in doc.tables:
            if len(table.rows) > 0:
                if "city" in table.rows[1].cells[0].text.lower() if len(table.rows) > 1 else False:
                    for row in table.rows:
                        row_text = row.cells[0].text.upper()
                        
                        if ("MADINAH" in row_text or "MEDINAH" in row_text) and madinah_hotel and "Select a" not in madinah_hotel:
                            self._write_bold_centered_cell(row.cells[1], madinah_hotel)
                            
                        if "MAKKAH" in row_text and makkah_hotel and "Select a" not in makkah_hotel:
                            self._write_bold_centered_cell(row.cells[1], makkah_hotel)

        # Ensure safe filename
        safe_name = "".join([c for c in data.passenger_name if c.isalpha() or c.isspace()]).rstrip()
        filename = f"{safe_name.replace(' ', '_')}_Itinerary.docx"
        output_path = os.path.join(self.output_dir, filename)
        
        doc.save(output_path)
        logger.info(f"DOCX saved successfully at {output_path}")
        
        return output_path
